# -*- coding: utf-8 -*-
"""
================================================================================
阶段三：问题2 —— 确定性备餐优化模型（MILP）
================================================================================

针对 2026-06-29，利用 PuLP + CBC 求解混合整数线性规划（MILP），在
「后厨人工 / 主灶设备 / 三大原材料」等物理约束下，求午、晚餐 10 种菜品
的最优生产批次数与备餐量，使综合经营效益全局最优，并将结果回填至
《附件4_结果汇总表.xlsx》的 Q2（及 Q1）Sheet。

---- 数学模型 ----
变量：
  x[j,m]  —— 菜品 j 在时段 m 的标准生产批次数（整数，≥0）
  y[j,m]  —— 备餐量 = B_j * x[j,m]（线性表达式，份）
  w[j,m]  —— 实际销量（连续，0 ≤ w ≤ min(y, D)）

目标（综合经营效益最大化）：
  max Σ Σ [ (S_j+L_j+P_j)·w[j,m] - (C_j+L_j)·y[j,m] - P_j·D[j,m] ]
  由于 (S_j+L_j+P_j) > 0，极大化会自然把 w 推向其上限 min(y, D)，实现完美线性化。

约束：
  1) w[j,m] ≤ y[j,m]                          （销量不超过备餐量）
  2) w[j,m] ≤ D[j,m]                          （销量不超过需求预测）
  3) y[j,m] ≤ MaxCapacity[j]                  （单时段最大生产能力，如有）
  4) 人工：Σ_j labor[j]·x[j,m] ≤ Labor[m]     （时段独立 / 或全天累加，自适应）
  5) 主灶：Σ_j stove[j]·x[j,m] ≤ Stove[m]
  6) 原料：Σ_{j,m} 用量[j]·x[j,m] ≤ 库存       （全天累加 / 或时段独立，自适应）

---- 自适应资源解析（工程亮点）----
  附件3 的《时段生产资源》按午/晚餐分别给出人工、主灶上限 => 时段独立约束；
  《当日原料供应》给出肉禽/蔬菜/蛋豆的全天总量 => 全天累加约束。
  脚本会自动识别“分时段”还是“全天总量”，无需手工改代码。

运行示例：
  python optimize_deterministic.py
================================================================================
"""

from __future__ import annotations

import argparse
import json
import logging
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pulp
import openpyxl
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# 全局常量
# ---------------------------------------------------------------------------
LOG = logging.getLogger("optimize")

# 输入文件
PREDICTIONS_CSV = "predictions_20260629.csv"
ATTACH3 = "附件3_菜品生产经营参数.xlsx"
ATTACH4 = "附件4_结果汇总表.xlsx"

# 附件3 的 sheet 名（优先匹配，缺失时回退到含关键字的 sheet）
SHEET_BASE_PARAM = "菜品基本经营参数"
SHEET_BATCH_PARAM = "批量生产参数"
SHEET_MATERIAL_USAGE = "主要原料消耗"
SHEET_TIME_RESOURCE = "时段生产资源"
SHEET_DAILY_SUPPLY = "当日原料供应"

# 时段 / 原料 / 菜品
MEALS = ["午餐", "晚餐"]
MATERIALS = ["肉禽类", "蔬菜类", "蛋豆类"]
DISH_PREFIX = "D"

# 求解器与容差
SOLVER_MSG = 0          # CBC 输出信息量（0=静默，1=详细）
BOTTLENECK_TOL = 0.001  # 占用率 >= 99.9% 视为瓶颈（饱和）
NEAR_TOL = 0.95         # 占用率 >= 95% 视为接近瓶颈

pd.set_option("display.width", 220)
pd.set_option("display.max_columns", None)


# ===========================================================================
# 通用工具
# ===========================================================================
def banner(msg: str) -> None:
    LOG.info("")
    LOG.info("=" * 92)
    LOG.info(msg)
    LOG.info("=" * 92)


def setup_logging(output_dir: Path, verbose: bool) -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError, OSError):
            pass
    level = logging.DEBUG if verbose else logging.INFO
    fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", datefmt="%H:%M:%S")
    root = logging.getLogger()
    root.setLevel(level)
    if not any(isinstance(h, logging.StreamHandler) for h in root.handlers):
        sh = logging.StreamHandler(sys.stdout)
        sh.setFormatter(fmt)
        root.addHandler(sh)
    output_dir.mkdir(parents=True, exist_ok=True)
    fh = logging.FileHandler(output_dir / "optimize_deterministic.log",
                             encoding="utf-8", mode="w")
    fh.setFormatter(fmt)
    root.addHandler(fh)
    LOG.setLevel(level)


def to_markdown(df: pd.DataFrame) -> str:
    """将 DataFrame 渲染为 GitHub 风格 Markdown 表格（用于精美控制台输出）。"""
    cols = list(df.columns)
    lines = ["| " + " | ".join(str(c) for c in cols) + " |",
             "|" + "|".join(" --- " for _ in cols) + "|"]
    for _, row in df.iterrows():
        cells = []
        for c in cols:
            v = row[c]
            if isinstance(v, float):
                v = f"{v:.2f}"
            cells.append("" if pd.isna(v) else str(v))
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def find_sheet(xls_sheets: list[str], hint: str, keywords: list[str]) -> str:
    """优先匹配 hint，否则匹配含任一关键字的 sheet。"""
    if hint in xls_sheets:
        return hint
    for s in xls_sheets:
        if any(k in s for k in keywords):
            return s
    raise KeyError(f"附件3 中找不到「{hint}」相关 Sheet，现有：{xls_sheets}")


def find_col(df: pd.DataFrame, *keywords: str) -> str:
    """按关键字模糊匹配列名（所有关键字必须同时命中）。"""
    for c in df.columns:
        name = str(c)
        if all(k in name for k in keywords):
            return c
    raise KeyError(f"找不到包含 {keywords} 的列，现有列：{list(df.columns)}")


def write_csv(df: pd.DataFrame, name: str, out_dir: Path, mirror_dir: Path | None) -> None:
    p = out_dir / name
    df.to_csv(p, index=False, encoding="utf-8-sig")
    LOG.info("已写出 %s（%d 行 × %d 列）", p.resolve(), len(df), df.shape[1])
    if mirror_dir is not None:
        mp = mirror_dir / name
        mirror_dir.mkdir(parents=True, exist_ok=True)
        df.to_csv(mp, index=False, encoding="utf-8-sig")
        LOG.info("镜像写出 %s", mp.resolve())


# ===========================================================================
# 数据加载（自适应解析）
# ===========================================================================
def load_predictions(path: Path) -> dict[tuple[str, str], float]:
    """读取阶段二预测：{(时段, 菜品编号): 预测需求量}。默认取「预测需求量」整数列。"""
    df = pd.read_csv(path)
    key_col = "预测需求量" if "预测需求量" in df.columns else "预测值_未取整"
    LOG.info("读取需求预测：%s（%d 行），采用字段「%s」",
             path.name, len(df), key_col)
    demand = {}
    for _, r in df.iterrows():
        meal = str(r["时段"]).strip()
        dish = str(r["菜品编号"]).strip()
        demand[(meal, dish)] = float(r[key_col])
    return demand


def _read_sheet_df(att3: Path, hint: str, keywords: list[str]) -> pd.DataFrame:
    xls = pd.ExcelFile(att3)
    sheet = find_sheet(xls.sheet_names, hint, keywords)
    df = pd.read_excel(att3, sheet_name=sheet)
    LOG.info("解析 Sheet「%s」（shape=%s）", sheet, df.shape)
    return df


def load_dish_params(att3: Path) -> dict[str, dict]:
    """
    读取菜品生产经营参数：S(售价), C(单位生产成本), L(剩余处理损失), P(缺供折算损失)；
    以及批量参数 B、人工耗时、主灶耗时、单时段最大产量；原料消耗。
    返回：
      {
        dish: {name, S, C, L, P, B, labor, stove, max_cap, meat, veg, eggbean}
      }
    """
    base = _read_sheet_df(att3, SHEET_BASE_PARAM, ["基本经营", "售价"])
    batch = _read_sheet_df(att3, SHEET_BATCH_PARAM, ["批量", "批次份数"])
    usage = _read_sheet_df(att3, SHEET_MATERIAL_USAGE, ["原料消耗", "消耗"])

    c_id = find_col(base, "菜品编号")
    c_name = find_col(base, "菜品名称")
    c_S = find_col(base, "售价")
    c_C = find_col(base, "单位生产成本")
    c_L = find_col(base, "剩余处理损失")
    c_P = find_col(base, "缺供")

    b_id = find_col(batch, "菜品编号")
    b_B = find_col(batch, "批次份数")
    b_labor = find_col(batch, "人工")
    b_stove = find_col(batch, "主灶")
    b_maxcap = find_col(batch, "最大产量")

    u_id = find_col(usage, "菜品编号")
    u_meat = find_col(usage, "肉禽")
    u_veg = find_col(usage, "蔬菜")
    u_egg = find_col(usage, "蛋豆")

    params: dict[str, dict] = {}
    for _, r in base.iterrows():
        dish = str(r[c_id]).strip()
        params[dish] = {
            "name": str(r[c_name]).strip(),
            "S": float(r[c_S]), "C": float(r[c_C]),
            "L": float(r[c_L]), "P": float(r[c_P]),
        }
    for _, r in batch.iterrows():
        dish = str(r[b_id]).strip()
        params[dish].update({
            "B": int(r[b_B]), "labor": float(r[b_labor]),
            "stove": float(r[b_stove]), "max_cap": float(r[b_maxcap]),
        })
    for _, r in usage.iterrows():
        dish = str(r[u_id]).strip()
        params[dish].update({
            "meat": float(r[u_meat]), "veg": float(r[u_veg]), "eggbean": float(r[u_egg]),
        })
    LOG.info("已解析 %d 个菜品的生产经营参数", len(params))
    return params


def load_resources(att3: Path) -> dict:
    """
    自适应读取生产资源：
      - 时段生产资源：若按午/晚餐分行 -> 时段独立约束；若仅全天单行 -> 全天累加。
      - 当日原料供应：若含午/晚餐列 -> 时段独立；若仅给出全天总量 -> 全天累加。
    返回：
      {
        "labor":  dict（时段->上限 或 {'全天':上限}）,
        "stove":  dict（同上）,
        "material": dict（原料->全天上限，或 时段->{原料:上限}）,
        "labor_per_meal": bool, "stove_per_meal": bool, "material_per_meal": bool,
      }
    """
    tres = _read_sheet_df(att3, SHEET_TIME_RESOURCE, ["生产资源", "时段"])
    t_meal = find_col(tres, "时段")
    t_labor = find_col(tres, "人工", "可用量")
    t_stove = find_col(tres, "主灶", "可用量")

    meal_rows = tres[[t_meal, t_labor, t_stove]].copy()
    meal_rows[t_meal] = meal_rows[t_meal].astype(str).str.strip()
    has_meal = meal_rows[t_meal].isin(MEALS).any()

    if has_meal:
        labor = {str(r[t_meal]): float(r[t_labor]) for _, r in meal_rows.iterrows() if str(r[t_meal]) in MEALS}
        stove = {str(r[t_meal]): float(r[t_stove]) for _, r in meal_rows.iterrows() if str(r[t_meal]) in MEALS}
        labor_per_meal = stove_per_meal = True
        LOG.info("时段生产资源：检测为【分时段】→ 建立时段独立约束：人工%s 主灶%s", labor, stove)
    else:
        labor = {"全天": float(meal_rows[t_labor].sum())}
        stove = {"全天": float(meal_rows[t_stove].sum())}
        labor_per_meal = stove_per_meal = False
        LOG.info("时段生产资源：检测为【全天总量】→ 建立全天累加约束：人工%s 主灶%s", labor, stove)

    sup = _read_sheet_df(att3, SHEET_DAILY_SUPPLY, ["原料供应", "可用量"])
    # 检测是否按午/晚餐分列
    sup_cols = [str(c) for c in sup.columns]
    meal_split = any(m in c for c in sup_cols for m in MEALS)
    mat_col = find_col(sup, "原料")
    amt_col = find_col(sup, "可用量")

    material: dict = {}
    if meal_split:
        material = {str(r[mat_col]).strip(): None for _, r in sup.iterrows()}
        for _, r in sup.iterrows():
            mat = str(r[mat_col]).strip()
            material[mat] = {m: float(r[find_col(sup, m)]) for m in MEALS}
        material_per_meal = True
        LOG.info("当日原料供应：检测为【分时段】→ 时段独立原料约束")
    else:
        for _, r in sup.iterrows():
            mat = str(r[mat_col]).strip()
            material[mat] = float(r[amt_col])
        material_per_meal = False
        LOG.info("当日原料供应：检测为【全天总量】→ 全天累加原料约束：%s", material)

    return {
        "labor": labor, "stove": stove, "material": material,
        "labor_per_meal": labor_per_meal, "stove_per_meal": stove_per_meal,
        "material_per_meal": material_per_meal,
    }


# ===========================================================================
# MILP 建模与求解
# ===========================================================================
def build_and_solve(params: dict, demand: dict, resources: dict) -> dict:
    dishes = sorted(params.keys())
    meals = MEALS

    prob = pulp.LpProblem("Deterministic_Prep_Optimization", pulp.LpMaximize)

    # ---- 变量 ----
    x = pulp.LpVariable.dicts("x", (dishes, meals), lowBound=0, cat="Integer")
    w = pulp.LpVariable.dicts("w", (dishes, meals), lowBound=0, cat="Continuous")
    # y = 备餐量（份），作为线性表达式 B*x
    y = {}
    for j in dishes:
        for m in meals:
            y[(j, m)] = params[j]["B"] * x[j][m]

    # ---- 目标函数 ----
    obj_terms = []
    for j in dishes:
        for m in meals:
            D = demand.get((m, j), 0.0)
            S, C, L, P = params[j]["S"], params[j]["C"], params[j]["L"], params[j]["P"]
            obj_terms.append((S + L + P) * w[j][m] - (C + L) * y[(j, m)] - P * D)
    prob += pulp.lpSum(obj_terms), "Total_Benefit"

    # ---- 约束 ----
    for j in dishes:
        for m in meals:
            D = demand.get((m, j), 0.0)
            # 1) 销量 <= 备餐量
            prob += w[j][m] <= y[(j, m)], f"sales_le_prep_{j}_{m}"
            # 2) 销量 <= 需求预测
            prob += w[j][m] <= D, f"sales_le_demand_{j}_{m}"
            # 3) 单时段最大生产能力
            prob += y[(j, m)] <= params[j]["max_cap"], f"maxcap_{j}_{m}"

    # ---- 人工 / 主灶约束（自适应：时段独立 or 全天累加）----
    labor = resources["labor"]
    stove = resources["stove"]
    if resources["labor_per_meal"]:
        for m in meals:
            prob += (pulp.lpSum(params[j]["labor"] * x[j][m] for j in dishes)
                     <= labor.get(m, 0.0), f"labor_{m}")
    else:
        prob += (pulp.lpSum(params[j]["labor"] * x[j][m] for j in dishes for m in meals)
                 <= labor.get("全天", 0.0), "labor_daily")

    if resources["stove_per_meal"]:
        for m in meals:
            prob += (pulp.lpSum(params[j]["stove"] * x[j][m] for j in dishes)
                     <= stove.get(m, 0.0), f"stove_{m}")
    else:
        prob += (pulp.lpSum(params[j]["stove"] * x[j][m] for j in dishes for m in meals)
                 <= stove.get("全天", 0.0), "stove_daily")

    # ---- 原料约束（自适应）----
    material = resources["material"]
    mat_cols = {"肉禽类": "meat", "蔬菜类": "veg", "蛋豆类": "eggbean"}
    for mat, field in mat_cols.items():
        if resources["material_per_meal"]:
            for m in meals:
                avail = material.get(mat, {}).get(m, 0.0)
                prob += (pulp.lpSum(params[j][field] * x[j][m] for j in dishes)
                         <= avail, f"{field}_{m}")
        else:
            avail = material.get(mat, 0.0)
            prob += (pulp.lpSum(params[j][field] * x[j][m] for j in dishes for m in meals)
                     <= avail, f"{field}_daily")

    n_int = len(dishes) * len(meals)
    n_cont = len(dishes) * len(meals)
    LOG.info("MILP 模型规模：%d 个整数变量 + %d 个连续变量，%d 个约束",
             n_int, n_cont, prob.numConstraints())

    # ---- 求解 ----
    solver = pulp.PULP_CBC_CMD(msg=SOLVER_MSG, timeLimit=120)
    prob.solve(solver)
    status = pulp.LpStatus[prob.status]
    LOG.info("求解状态：%s", status)

    if prob.status != pulp.LpStatusOptimal:
        LOG.error("未获得最优解（状态=%s）。请检查约束是否可行。", status)
        raise RuntimeError(f"MILP 求解状态非最优：{status}")

    # ---- 提取解 ----
    x_opt = {j: {m: int(round(x[j][m].varValue)) for m in meals} for j in dishes}
    y_opt = {j: {m: x_opt[j][m] * params[j]["B"] for m in meals} for j in dishes}
    w_opt = {j: {m: min(y_opt[j][m], demand.get((m, j), 0.0)) for m in meals} for j in dishes}

    return {
        "prob": prob, "dishes": dishes, "meals": meals,
        "x": x_opt, "y": y_opt, "w": w_opt,
        "objective": float(pulp.value(prob.objective)),
        "status": status,
    }


# ===========================================================================
# 后评估与瓶颈分析
# ===========================================================================
def evaluate(solution: dict, params: dict, demand: dict, resources: dict) -> dict:
    dishes = solution["dishes"]
    meals = solution["meals"]
    x, y, w = solution["x"], solution["y"], solution["w"]

    # ---- 方案明细 ----
    plan_rows = []
    for m in meals:
        for j in dishes:
            D = demand.get((m, j), 0.0)
            plan_rows.append({
                "时段": m, "菜品编号": j, "菜品名称": params[j]["name"],
                "预测需求量": int(round(D)),
                "生产批次数(批)": int(x[j][m]),
                "备餐量(份)": int(y[j][m]),
                "预计销量(份)": int(round(w[j][m])),
                "预计剩余(份)": int(round(y[j][m] - w[j][m])),
                "预计缺供(份)": int(round(D - w[j][m])),
            })
    plan_df = pd.DataFrame(plan_rows)

    # ---- 经营指标 ----
    revenue = sum(params[j]["S"] * w[j][m] for j in dishes for m in meals)
    prod_cost = sum(params[j]["C"] * y[j][m] for j in dishes for m in meals)
    leftover_units = sum(y[j][m] - w[j][m] for j in dishes for m in meals)
    shortage_units = sum(demand.get((m, j), 0.0) - w[j][m] for j in dishes for m in meals)
    leftover_loss = sum(params[j]["L"] * (y[j][m] - w[j][m]) for j in dishes for m in meals)
    shortage_loss = sum(params[j]["P"] * (demand.get((m, j), 0.0) - w[j][m])
                        for j in dishes for m in meals)
    total_prepared = sum(y[j][m] for j in dishes for m in meals)
    total_sales = sum(w[j][m] for j in dishes for m in meals)
    benefit = revenue - prod_cost - leftover_loss - shortage_loss
    service_level = total_sales / sum(demand.get((m, j), 0.0) for j in dishes for m in meals) * 100.0

    metrics = {
        "销售收入(元)": revenue, "生产成本(元)": prod_cost,
        "剩余损失(元)": leftover_loss, "缺供损失(元)": shortage_loss,
        "综合经营效益(元)": benefit,
        "备餐总量(份)": total_prepared, "预计销量(份)": total_sales,
        "预计剩余量(份)": leftover_units, "预计缺供量(份)": shortage_units,
        "服务水平(%)": service_level,
    }

    # ---- 资源占用率 ----
    res_rows = []
    labor = resources["labor"]
    stove = resources["stove"]
    material = resources["material"]

    def _used_labor(m):
        return sum(params[j]["labor"] * x[j][m] for j in dishes)

    def _used_stove(m):
        return sum(params[j]["stove"] * x[j][m] for j in dishes)

    if resources["labor_per_meal"]:
        for m in meals:
            u = _used_labor(m); a = labor.get(m, 0.0)
            res_rows.append({"资源": f"{m}人工", "实际使用量": round(u, 2),
                             "可用量": a, "利用率(%)": round(u / a * 100, 2) if a else 0.0})
    else:
        u = sum(_used_labor(m) for m in meals); a = labor.get("全天", 0.0)
        res_rows.append({"资源": "全天人工", "实际使用量": round(u, 2),
                         "可用量": a, "利用率(%)": round(u / a * 100, 2) if a else 0.0})

    if resources["stove_per_meal"]:
        for m in meals:
            u = _used_stove(m); a = stove.get(m, 0.0)
            res_rows.append({"资源": f"{m}主灶", "实际使用量": round(u, 2),
                             "可用量": a, "利用率(%)": round(u / a * 100, 2) if a else 0.0})
    else:
        u = sum(_used_stove(m) for m in meals); a = stove.get("全天", 0.0)
        res_rows.append({"资源": "全天主灶", "实际使用量": round(u, 2),
                         "可用量": a, "利用率(%)": round(u / a * 100, 2) if a else 0.0})

    mat_cols = {"肉禽类": "meat", "蔬菜类": "veg", "蛋豆类": "eggbean"}
    for mat, field in mat_cols.items():
        if resources["material_per_meal"]:
            for m in meals:
                u = sum(params[j][field] * x[j][m] for j in dishes)
                a = material.get(mat, {}).get(m, 0.0)
                res_rows.append({"资源": f"{m}{mat}", "实际使用量": round(u, 3),
                                 "可用量": a, "利用率(%)": round(u / a * 100, 2) if a else 0.0})
        else:
            u = sum(params[j][field] * x[j][m] for j in dishes for m in meals)
            a = material.get(mat, 0.0)
            res_rows.append({"资源": f"{mat}(全天)", "实际使用量": round(u, 3),
                             "可用量": a, "利用率(%)": round(u / a * 100, 2) if a else 0.0})

    res_df = pd.DataFrame(res_rows).sort_values("利用率(%)", ascending=False).reset_index(drop=True)

    # ---- 瓶颈识别 ----
    bottlenecks = res_df[res_df["利用率(%)"] >= (100.0 - BOTTLENECK_TOL * 100)].copy()
    near = res_df[(res_df["利用率(%)"] < (100.0 - BOTTLENECK_TOL * 100))
                  & (res_df["利用率(%)"] >= NEAR_TOL * 100)].copy()

    return {
        "plan_df": plan_df, "res_df": res_df, "metrics": metrics,
        "bottlenecks": bottlenecks, "near": near,
    }


# ===========================================================================
# Excel 回填
# ===========================================================================
def backfill_excel(att4: Path, out_dir: Path, mirror_dir: Path | None,
                   solution: dict, demand: dict, params: dict, eval_res: dict) -> None:
    """
    将最优方案回填至《附件4_结果汇总表.xlsx》：
      - Q1：D 列填预测需求量（问题1结果）
      - Q2：D6:I25 填批次数/备餐量/需求/销量/剩余/缺供；B30:C36 与 G30:G34 填资源与经营结果
    仅写业务单元格，保留其余 Sheet、公式与格式。
    """
    banner("任务3｜Excel 结果规范回填")

    # 1) 原始备份（仅首次），保证原始模板永不丢失
    backup = out_dir / "附件4_结果汇总表_原始备份.xlsx"
    if not backup.exists():
        shutil.copy2(att4, backup)
        LOG.info("已创建原始模板备份：%s", backup.resolve())

    wb = openpyxl.load_workbook(att4)  # data_only=False，保留公式
    meals = solution["meals"]
    dishes = solution["dishes"]

    # ---- Q1：预测需求量 ----
    if "Q1" in wb.sheetnames:
        ws1 = wb["Q1"]
        row = 6
        for m in meals:
            for j in dishes:
                ws1.cell(row=row, column=4).value = int(round(demand.get((m, j), 0.0)))
                row += 1
        LOG.info("Q1：已回填预测需求量（D6:D%d）", row - 1)

    # ---- Q2：方案明细 + 资源/经营结果 ----
    ws2 = wb["Q2"]
    x, y, w = solution["x"], solution["y"], solution["w"]

    # 方案明细（行 6-25）
    row = 6
    for m in meals:
        for j in dishes:
            D = demand.get((m, j), 0.0)
            ws2.cell(row=row, column=4).value = int(x[j][m])                      # 生产批次数
            ws2.cell(row=row, column=5).value = int(y[j][m])                      # 备餐量
            ws2.cell(row=row, column=6).value = int(round(D))                     # 需求估计量
            ws2.cell(row=row, column=7).value = int(round(w[j][m]))               # 预计销量
            ws2.cell(row=row, column=8).value = int(round(y[j][m] - w[j][m]))     # 预计剩余
            ws2.cell(row=row, column=9).value = int(round(D - w[j][m]))           # 预计缺供
            row += 1
    LOG.info("Q2：已回填 20 条方案明细（D6:I%d）", row - 1)

    # 资源占用（行 30-36）与经营结果（G30:G34）
    resources = eval_res["res_df"].set_index("资源")
    metrics = eval_res["metrics"]

    def _res(key_contains: str) -> dict:
        for name, r in resources.iterrows():
            if key_contains in str(name):
                return r
        return {"实际使用量": 0.0, "可用量": 0.0}

    summary_map = [
        (30, "午餐人工"), (31, "午餐主灶"), (32, "晚餐人工"), (33, "晚餐主灶"),
        (34, "肉禽类"), (35, "蔬菜类"), (36, "蛋豆类"),
    ]
    for r, label in summary_map:
        rrow = _res(label)
        ws2.cell(row=r, column=2).value = round(float(rrow["实际使用量"]), 3)   # 实际使用量
        ws2.cell(row=r, column=3).value = round(float(rrow["可用量"]), 3)       # 可用量

    # 经营结果汇总（G 列）
    result_map = {
        30: metrics["综合经营效益(元)"],
        31: metrics["备餐总量(份)"],
        32: metrics["预计销量(份)"],
        33: metrics["预计剩余量(份)"],
        34: metrics["预计缺供量(份)"],
    }
    for r, v in result_map.items():
        ws2.cell(row=r, column=7).value = round(float(v), 2)
    LOG.info("Q2：已回填资源占用（B30:C36）与经营结果（G30:G34），保留 D 列利用率公式")

    # ---- 保存 ----
    wb.save(att4)
    LOG.info("已回填并保存：%s", att4.resolve())

    # 另存一份已回填副本到输出目录
    filled_copy = out_dir / "附件4_结果汇总表_已回填.xlsx"
    wb.save(filled_copy)
    LOG.info("已回填副本：%s", filled_copy.resolve())
    if mirror_dir is not None:
        mirror_dir.mkdir(parents=True, exist_ok=True)
        wb.save(mirror_dir / "附件4_结果汇总表_已回填.xlsx")
        LOG.info("镜像副本：%s", (mirror_dir / "附件4_结果汇总表_已回填.xlsx").resolve())


# ===========================================================================
# 主流程
# ===========================================================================
def run_pipeline(args: argparse.Namespace) -> None:
    data_dir = Path(args.data_dir)
    out_dir = Path(args.output_dir)
    mirror_dir = Path(args.mirror_dir) if args.mirror_dir else None
    for d in (data_dir, out_dir):
        d.mkdir(parents=True, exist_ok=True)
    if mirror_dir is not None:
        mirror_dir.mkdir(parents=True, exist_ok=True)
    setup_logging(out_dir, args.verbose)

    banner("阶段三｜问题2 —— 确定性备餐优化（MILP）启动")
    LOG.info("数据目录：%s；输出目录：%s", data_dir.resolve(), out_dir.resolve())

    # ---- 数据加载 ----
    # 预测文件可能在 D:\B\output\ 或镜像目录，多候选稳健定位
    pred_candidates = [data_dir / PREDICTIONS_CSV, data_dir / "output" / PREDICTIONS_CSV]
    if mirror_dir is not None:
        pred_candidates.append(mirror_dir / PREDICTIONS_CSV)
    pred_path = next((p for p in pred_candidates if p.exists()), None)
    if pred_path is None:
        raise FileNotFoundError(f"找不到 {PREDICTIONS_CSV}，已搜索：{pred_candidates}")
    att3 = data_dir / ATTACH3
    att4 = data_dir / ATTACH4
    for f in (att3, att4):
        if not f.exists():
            raise FileNotFoundError(f"找不到输入文件：{f}")

    demand = load_predictions(pred_path)
    params = load_dish_params(att3)
    resources = load_resources(att3)

    # ---- 求解 ----
    solution = build_and_solve(params, demand, resources)
    LOG.info("MILP 全局最优综合经营效益（目标值）：%.2f", solution["objective"])

    # ---- 后评估 ----
    eval_res = evaluate(solution, params, demand, resources)

    banner("任务1｜最优备餐方案")
    LOG.info("\n%s", to_markdown(eval_res["plan_df"]))

    banner("任务2｜资源占用率与瓶颈分析")
    LOG.info("\n%s", to_markdown(eval_res["res_df"]))
    if len(eval_res["bottlenecks"]):
        LOG.warning("【瓶颈约束】占用率达到 100% 的资源：\n%s",
                    to_markdown(eval_res["bottlenecks"]))
    if len(eval_res["near"]):
        LOG.info("【接近瓶颈】占用率 ≥ %.0f%% 的资源：\n%s",
                 NEAR_TOL * 100, to_markdown(eval_res["near"]))
    if not len(eval_res["bottlenecks"]) and not len(eval_res["near"]):
        LOG.info("无瓶颈约束：所有资源均有富余。")

    banner("任务2｜业绩指标预测")
    metrics_df = pd.DataFrame(
        [{"指标": k, "数值": round(v, 2), "单位": "元" if "元" in k else
          ("%" if "%" in k else "份")} for k, v in eval_res["metrics"].items()])
    LOG.info("\n%s", to_markdown(metrics_df))
    LOG.info("（模型目标值与后评估效益应一致：目标=%.2f，后评估=%.2f）",
             solution["objective"], eval_res["metrics"]["综合经营效益(元)"])

    # ---- Excel 回填 ----
    backfill_excel(att4, out_dir, mirror_dir, solution, demand, params, eval_res)

    # ---- 导出 CSV ----
    banner("任务3｜导出结果 CSV")
    write_csv(eval_res["plan_df"], "q2_optimal_plan.csv", out_dir, mirror_dir)
    write_csv(metrics_df, "q2_performance_metrics.csv", out_dir, mirror_dir)
    write_csv(eval_res["res_df"], "q2_resource_utilization.csv", out_dir, mirror_dir)

    banner("阶段三完成")
    LOG.info("所有结果已写入：%s", out_dir.resolve())


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="阶段三：问题2 确定性备餐优化（MILP）")
    # 默认以脚本所在目录为基准，附件3/4 在脚本目录、输出到 脚本目录/output，换机无需改代码
    script_dir = Path(__file__).resolve().parent
    p.add_argument("--data-dir", default=str(script_dir), help="数据目录（含附件3/4，默认：脚本目录）")
    p.add_argument("--output-dir", default=str(script_dir / "output"), help="输出目录（默认：脚本目录/output）")
    p.add_argument("--mirror-dir", default="", help="镜像目录（默认禁用；如需双写可指定路径）")
    p.add_argument("--verbose", action="store_true", help="DEBUG 日志")
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    try:
        run_pipeline(args)
    except Exception as exc:
        logging.getLogger().exception("流水线执行失败：%s", exc)
        sys.exit(1)


if __name__ == "__main__":
    main()
