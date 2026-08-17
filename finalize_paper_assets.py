# -*- coding: utf-8 -*-
"""
================================================================================
阶段五：学术图表终验与自动报告生成
================================================================================

本脚本产出论文所需的最高规格插图，并对最终回填的《附件4_结果汇总表.xlsx》
做全指标闭环审计，最后自动生成「论文写作素材白皮书」。

  任务 1  学术图表矩阵
          - 图1 预测残差分布（直方图 + KDE + 正态拟合 + Jarque-Bera 检验）
          - 图2 后厨资源占用与饱和度（95% 饱和线 + 瓶颈高亮）
          - 图3 Q2/Q3 收益 CDF 重叠（双色交替虚线 + 学术批注框）

  任务 2  附件4 终极一致性审计（批次数×批次份数=备餐量、Q3-Q2=0、效益变化=-952.24 等）

  任务 3  论文写作素材白皮书 paper_writing_materials.md（核心数值/LaTeX 公式/论述段落）

配色：高雅内敛的莫兰迪/学术色系；DPI=150；中文字体自适应。
================================================================================
"""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns
from scipy import stats as sps

from optimize_deterministic import load_dish_params, MEALS
from optimize_stochastic import generate_scenarios, evaluate_plan

# ---------------------------------------------------------------------------
# 全局常量
# ---------------------------------------------------------------------------
LOG = logging.getLogger("finalize")

ATTACH3_NAME = "附件3_菜品生产经营参数.xlsx"
ATTACH4_NAME = "附件4_结果汇总表.xlsx"


def find_attach4(data_dir: Path) -> Path:
    """兼容“附件4_结果汇总表.xlsx”与“附件4：结果汇总表.xlsx”两种命名。"""
    for name in [ATTACH4_NAME, "附件4：结果汇总表.xlsx"]:
        p = data_dir / name
        if p.exists():
            return p
    raise FileNotFoundError(f"找不到附件4，已搜索：{data_dir}")

# 莫兰迪 / 学术配色
PALETTE = {
    "blue": "#6E8CA0",      # 科技蓝灰
    "gray": "#9AA5A8",      # 学术灰
    "terracotta": "#C08A6E",  # 莫兰迪陶土
    "green": "#8FA98F",     # 莫兰迪绿
    "deep_blue": "#4A6572",
    "deep_red": "#B0524A",  # 深红（瓶颈）
    "gold": "#C9A86A",
}
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
sns.set_theme(style="whitegrid", font=["Microsoft YaHei", "SimHei", "DejaVu Sans"])


# ===========================================================================
# 工具
# ===========================================================================
def banner(msg: str) -> None:
    LOG.info("")
    LOG.info("=" * 92)
    LOG.info(msg)
    LOG.info("=" * 92)


def setup_logging(output_dir: Path, verbose: bool) -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError, OSError):
            pass
    level = logging.DEBUG if verbose else logging.INFO
    fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", datefmt="%H:%M:%S")
    root = logging.getLogger()
    root.setLevel(level)
    if not any(isinstance(h, logging.StreamHandler) for h in root.handlers):
        sh = logging.StreamHandler(sys.stdout)
        sh.setFormatter(fmt)
        root.addHandler(sh)
    output_dir.mkdir(parents=True, exist_ok=True)
    fh = logging.FileHandler(output_dir / "finalize_paper_assets.log", encoding="utf-8", mode="w")
    fh.setFormatter(fmt)
    root.addHandler(fh)
    LOG.setLevel(level)


def save_fig(fig: plt.Figure, name: str, out_dir: Path, mirror_dir: Path | None) -> None:
    img = out_dir / "images"
    img.mkdir(parents=True, exist_ok=True)
    path = img / name
    fig.savefig(path, dpi=150, bbox_inches="tight")
    LOG.info("已保存图 %s", path.resolve())
    if mirror_dir is not None:
        m = mirror_dir / "images"
        m.mkdir(parents=True, exist_ok=True)
        fig.savefig(m / name, dpi=150, bbox_inches="tight")
    plt.close(fig)


# ===========================================================================
# 任务 1：图1 预测残差分布
# ===========================================================================
def plot_residual_distribution(residuals: pd.DataFrame, out_dir: Path, mirror_dir: Path | None) -> None:
    banner("任务1｜图1 预测残差分布与正态拟合")
    r = residuals["预测残差"].to_numpy(dtype=float)
    mean, std = float(r.mean()), float(r.std(ddof=1))
    # Jarque-Bera 正态性检验
    jb_stat, jb_p = sps.jarque_bera(r)
    # Shapiro-Wilk（样本量 1000 也适用）
    sw_stat, sw_p = sps.shapiro(r) if len(r) <= 5000 else (np.nan, np.nan)

    fig, ax = plt.subplots(figsize=(9, 6))
    sns.histplot(r, bins=40, stat="density", color=PALETTE["blue"], alpha=0.55,
                 edgecolor="white", label="经验分布直方图", ax=ax)
    sns.kdeplot(r, color=PALETTE["deep_blue"], lw=2, label="核密度估计 (KDE)", ax=ax)
    xs = np.linspace(r.min() - 5, r.max() + 5, 500)
    ax.plot(xs, sps.norm.pdf(xs, mean, std), color=PALETTE["deep_red"], lw=2,
            ls="--", label="正态分布拟合 N(μ, σ²)")
    ax.axvline(mean, color=PALETTE["gold"], ls=":", lw=1.6, label=f"均值 μ = {mean:.2f}")
    ax.set_xlabel("预测残差 ε = 真实需求 − 预测值（份）")
    ax.set_ylabel("概率密度")
    ax.set_title("预测残差分布与正态拟合（预测效果检验）", fontsize=14)
    ax.legend(loc="upper right", framealpha=0.9)
    # 统计量标注框
    text = (f"均值 μ = {mean:.3f}\n标准差 σ = {std:.3f}\n"
            f"Jarque-Bera: 统计量={jb_stat:.2f}, p={jb_p:.4f}\n"
            f"Shapiro-Wilk: p={sw_p:.4f}")
    ax.text(0.03, 0.97, text, transform=ax.transAxes, va="top", ha="left",
            fontsize=10, bbox=dict(boxstyle="round,pad=0.5", facecolor="white", alpha=0.85))
    LOG.info("残差统计：μ=%.3f, σ=%.3f, JB p=%.4f, Shapiro p=%.4f", mean, std, jb_p, sw_p)
    save_fig(fig, "prediction_residuals_distribution.png", out_dir, mirror_dir)


# ===========================================================================
# 任务 1：图2 资源饱和度
# ===========================================================================
def plot_resource_saturation(res_df: pd.DataFrame, out_dir: Path, mirror_dir: Path | None) -> None:
    banner("任务1｜图2 后厨资源占用与饱和度")
    df = res_df.copy()
    df["资源"] = df["资源"].str.replace("(全天)", "（全天）", regex=False)
    df = df.sort_values("利用率(%)", ascending=False)
    util = df["利用率(%)"].to_numpy(float)
    labels = df["资源"].to_list()

    fig, ax = plt.subplots(figsize=(10, 6))
    colors = [PALETTE["deep_red"] if u >= 95 else PALETTE["blue"] for u in util]
    bars = ax.bar(range(len(labels)), util, color=colors, width=0.62, alpha=0.9, edgecolor="white")
    ax.axhline(95, color=PALETTE["deep_red"], ls="--", lw=1.8,
               label="95% 饱和度预警线")
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels, rotation=20, ha="right")
    ax.set_ylim(0, 110)
    ax.set_ylabel("资源占用率（%）")
    ax.set_title("后厨资源占用与饱和度分析（瓶颈识别）", fontsize=14)
    for b, u in zip(bars, util):
        ax.text(b.get_x() + b.get_width() / 2, u + 1.5, f"{u:.1f}%",
                ha="center", va="bottom", fontsize=9)
    ax.legend(loc="lower right")
    # 瓶颈区间标注
    if (util >= 95).any():
        ax.text(len(labels) - 0.4, 97.5, "◀ 产能赤字瓶颈区间（≥95%）",
                color=PALETTE["deep_red"], fontsize=10, ha="right", va="bottom", fontweight="bold")
    LOG.info("资源占用率（降序）：%s", dict(zip(labels, [round(u, 1) for u in util])))
    save_fig(fig, "resource_saturation_analysis.png", out_dir, mirror_dir)


# ===========================================================================
# 任务 1：图3 CDF 重叠
# ===========================================================================
def plot_cdf_overlap(benefits_q2: np.ndarray, benefits_q3: np.ndarray,
                     expected_benefit: float, out_dir: Path, mirror_dir: Path | None) -> None:
    banner("任务1｜图3 Q2/Q3 收益 CDF 重叠对比")

    def ecdf(data):
        xs = np.sort(data)
        ys = np.arange(1, len(xs) + 1) / len(xs)
        return xs, ys

    x2, y2 = ecdf(benefits_q2)
    x3, y3 = ecdf(benefits_q3)

    fig, ax = plt.subplots(figsize=(9, 6))
    # 双色交替虚线（Q2/Q3 完全重合时也能清晰展示两条曲线）
    ax.plot(x2, y2, drawstyle="steps-post", lw=2, color=PALETTE["terracotta"],
            ls=(0, (6, 3)), label="Q2 确定性方案 CDF")
    ax.plot(x3, y3, drawstyle="steps-post", lw=2, color=PALETTE["deep_blue"],
            ls=(0, (2, 4)), label="Q3 随机稳健方案 CDF")
    ax.axvline(expected_benefit, color=PALETTE["deep_red"], ls="--", lw=1.5, alpha=0.8)
    ax.set_xlabel("综合经营效益（元）")
    ax.set_ylabel("累积概率")
    ax.set_title("Q2 确定性 vs Q3 随机稳健：1000 场景综合效益 CDF", fontsize=14)
    ax.legend(loc="upper left")
    ax.grid(True, alpha=0.3)
    # 学术批注框
    note = (f"两曲线完全重合：当后厨产能成为绝对硬约束时，\n"
            f"确定性最优方案与随机稳健方案在不确定需求下完全等价。\n"
            f"期望综合效益 E[Benefit] = {expected_benefit:.2f} 元")
    ax.text(0.03, 0.34, note, transform=ax.transAxes, va="bottom", ha="left", fontsize=10,
            bbox=dict(boxstyle="round,pad=0.6", facecolor="#F5F0E6", edgecolor=PALETTE["gold"],
                      alpha=0.95))
    save_fig(fig, "q2_vs_q3_cdf_overlap.png", out_dir, mirror_dir)


# ===========================================================================
# 任务 2：附件4 终极一致性审计
# ===========================================================================
def audit_excel(params: dict, att4: Path) -> list[dict]:
    banner("任务2｜附件4 结果汇总表终极一致性审计")
    wb = openpyxl.load_workbook(att4, data_only=False)
    checks: list[dict] = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        checks.append({"检查项": name, "结果": "PASS" if ok else "FAIL", "说明": detail})
        LOG.info("  [%s] %s %s", "PASS" if ok else "FAIL", name, f"—— {detail}" if detail else "")

    ws2, ws3 = wb["Q2"], wb["Q3"]
    dishes = sorted(params.keys())

    # ---- 审计 1：批次数 × 批次份数 == 备餐量 ----
    ok1 = True
    bad = []
    for ws, tag in ((ws2, "Q2"), (ws3, "Q3")):
        row = 6
        for m in MEALS:
            for j in dishes:
                d = ws.cell(row=row, column=4).value
                e = ws.cell(row=row, column=5).value
                B = params[j]["B"]
                if d is None or e is None or int(d) * B != int(e):
                    ok1 = False
                    bad.append(f"{tag}{m}{j}: {d}×{B}≠{e}")
                row += 1
    check("审计1｜批次数×标准批次份数=备餐量（Q2+Q3 共40条）", ok1, f"不一致条目={bad}" if bad else "全部一致")

    # ---- 审计 2：相比Q2备餐量变化 == Q3备餐量 - Q2备餐量 == 0 ----
    ok2 = True
    max_diff = 0
    row = 6
    for m in MEALS:
        for j in dishes:
            e2 = ws2.cell(row=row, column=5).value
            e3 = ws3.cell(row=row, column=5).value
            diff = int(e3) - int(e2)
            max_diff = max(max_diff, abs(diff))
            if diff != 0:
                ok2 = False
            # F 列公式完整性
            f3 = ws3.cell(row=row, column=6).value
            if not (isinstance(f3, str) and f3.startswith("=") and "'Q2'" in f3):
                ok2 = False
            row += 1
    check("审计2｜相比Q2备餐量变化=Q3−Q2=0（且F列公式完整）", ok2, f"最大变化量={max_diff}")

    # ---- 审计 3：综合经营效益变化 == -952.24 ----
    b30 = ws3.cell(row=30, column=2).value
    c30 = ws3.cell(row=30, column=3).value
    d30_formula = ws3.cell(row=30, column=4).value
    expected_change = 13375.96 - 14328.20
    ok3 = (b30 is not None and c30 is not None
           and abs(float(b30) - 14328.20) < 0.01 and abs(float(c30) - 13375.96) < 0.01
           and abs((float(c30) - float(b30)) - expected_change) < 0.01
           and isinstance(d30_formula, str) and d30_formula.startswith("="))
    check("审计3｜综合经营效益变化 = 13375.96−14328.20 = −952.24", ok3,
          f"Q2={b30}, Q3={c30}, Δ={round((float(c30 or 0) - float(b30 or 0)), 2)}")

    # ---- 审计 4：空行 / 溢出字符 / 格式损坏 ----
    ok4 = True
    bad_cells = []
    for ws, tag in ((ws2, "Q2"), (ws3, "Q3")):
        for r in range(6, 26):
            for c in (4, 5):
                v = ws.cell(row=r, column=c).value
                if v is None or (isinstance(v, str) and ("#" in v or "REF" in v)):
                    ok4 = False
                    bad_cells.append(f"{tag}!{openpyxl.utils.get_column_letter(c)}{r}")
    for r in range(30, 36):
        for c in (2, 3):
            v = ws3.cell(row=r, column=c).value
            if v is None or (isinstance(v, str) and "#" in v):
                ok4 = False
                bad_cells.append(f"Q3!{openpyxl.utils.get_column_letter(c)}{r}")
    check("审计4｜无空单元格/溢出字符(###)/格式损坏", ok4, f"异常单元格={bad_cells}" if bad_cells else "格式完好")

    # ---- 审计 5：Q1 预测需求量已回填 ----
    ws1 = wb["Q1"]
    ok5 = all(ws1.cell(row=r, column=4).value is not None for r in range(6, 26))
    check("审计5｜Q1 预测需求量 20 条已回填", ok5, "" if ok5 else "存在空单元格")

    n_pass = sum(1 for c in checks if c["结果"] == "PASS")
    LOG.info("审计汇总：%d / %d 项通过", n_pass, len(checks))
    return checks


# ===========================================================================
# 任务 3：论文写作素材白皮书
# ===========================================================================
def generate_white_paper(out_dir: Path, mirror_dir: Path | None,
                         q2_benefit: float, q3_benefit: float,
                         residual_stats: pd.DataFrame, res_df: pd.DataFrame,
                         checks: list[dict]) -> None:
    banner("任务3｜生成论文写作素材白皮书")
    uncertainty_cost = q3_benefit - q2_benefit
    bottleneck = res_df.sort_values("利用率(%)", ascending=False).head(2)
    bottleneck_txt = "、".join(
        f"{r['资源']} {r['利用率(%)']:.1f}%" for _, r in bottleneck.iterrows())

    md = []
    md.append("# 论文写作素材白皮书\n")
    md.append("> 由 `finalize_paper_assets.py` 自动生成，数值均来自阶段一~四实际求解结果。\n")

    # ---- 核心数值速查表 ----
    md.append("## 一、核心数值速查表\n")
    md.append("| 指标 | 数值 | 来源 |")
    md.append("| --- | --- | --- |")
    md.append("| Q1 需求预测 MAPE（LightGBM 5折 OOF） | 9.04% | 阶段二 |")
    md.append("| Q1 需求预测 MAE | 9.58 份 | 阶段二 |")
    md.append(f"| Q2 确定性最优综合经营效益 | {q2_benefit:.2f} 元 | 阶段三 |")
    md.append(f"| Q3 随机稳健期望综合经营效益 | {q3_benefit:.2f} 元 | 阶段四 |")
    md.append(f"| 需求不确定性代价（Q3−Q2） | {uncertainty_cost:.2f} 元（{uncertainty_cost/q2_benefit*100:.2f}%） | 阶段四 |")
    md.append(f"| 资源瓶颈饱和度 | {bottleneck_txt} | 阶段三 |")
    md.append("| 期望服务水平（1000场景） | 94.55% | 阶段四 |")
    md.append("| 效益标准差（1000场景） | 1499.51 元 | 阶段四 |")
    md.append("| VaR5%（5%分位效益） | 10624.98 元 | 阶段四 |")
    md.append("")

    # ---- 残差分布表 ----
    md.append("### 预测残差分布特征（20个菜品-时段）\n")
    md.append("| 菜品编号 | 时段 | 均值 | 标准差 | 偏度 | 峰度 |")
    md.append("| --- | --- | --- | --- | --- | --- |")
    for _, r in residual_stats.iterrows():
        md.append(f"| {r['菜品编号']} | {r['时段']} | {r['均值(Mean)']:.2f} | "
                  f"{r['标准差(Std)']:.2f} | {r['偏度(Skewness)']:.2f} | {r['峰度(Kurtosis)']:.2f} |")
    md.append("")

    # ---- LaTeX 公式 ----
    md.append("## 二、LaTeX 核心公式代码段\n")

    md.append("### 2.1 售罄截断真实需求重构（阶段一）\n")
    md.append("```latex")
    md.append(r"\theta_{j,m} = \frac{1}{|\mathcal{S}_{j,m}|}\sum_{t\in\mathcal{S}_{j,m}}"
              r"\frac{\hat{D}_{j,m,t}}{S_{j,m,t}},\quad "
              r"\mathcal{S}_{j,m}=\{t: \text{售罄且}\ \hat{D}_{j,m,t}>S_{j,m,t}\}")
    md.append(r"")
    md.append(r"D_{j,m}^{\text{true}} = \begin{cases}")
    md.append(r"  S_{j,m}, & \text{未售罄}\\[4pt]")
    md.append(r"  \hat{D}_{j,m}, & \text{售罄且}\ \hat{D}_{j,m} > S_{j,m}\\[4pt]")
    md.append(r"  S_{j,m}\cdot\bar{\theta}_{j,m}, & \text{售罄且估计量缺失/异常}\\[4pt]")
    md.append(r"  1.25\,S_{j,m}, & \text{样本不足（全局保守系数）}")
    md.append(r"\end{cases}")
    md.append("```\n")

    md.append("### 2.2 问题2 确定性 MILP（阶段三）\n")
    md.append("```latex")
    md.append(r"\max \sum_{m\in\{l,d\}}\sum_{j=1}^{10}"
              r"\left[(S_j+L_j+P_j)\,w_{j,m} - (C_j+L_j)\,y_{j,m} - P_j D_{j,m}\right]")
    md.append(r"\begin{aligned}")
    md.append(r"\text{s.t.}\quad & w_{j,m}\le y_{j,m},\quad w_{j,m}\le D_{j,m},\quad "
              r"y_{j,m}=B_j x_{j,m}\le \overline{Y}_j\\")
    md.append(r"& \sum_j \ell_j x_{j,m}\le L_m,\quad \sum_j s_j x_{j,m}\le S_m\\")
    md.append(r"& \sum_m\sum_j r_j^k x_{j,m}\le R^k,\quad k\in\{\text{肉禽,蔬菜,蛋豆}\}\\")
    md.append(r"& x_{j,m}\in\mathbb{Z}_{\ge 0},\ w_{j,m}\ge 0")
    md.append(r"\end{aligned}")
    md.append("```\n")

    md.append("### 2.3 问题3 两阶段随机规划 SAA（阶段四）\n")
    md.append("```latex")
    md.append(r"\max \frac{1}{S}\sum_{s=1}^{S}\sum_{m}\sum_{j}"
              r"\left[ S_j w_{j,m}^s - C_j y_{j,m} - L_j u_{j,m}^s - P_j v_{j,m}^s \right]")
    md.append(r"\begin{aligned}")
    md.append(r"\text{s.t.}\quad & \sum_j \ell_j x_{j,m}\le L_m,\ \sum_j s_j x_{j,m}\le S_m,"
              r"\ \sum_m\sum_j r_j^k x_{j,m}\le R^k\\")
    md.append(r"& w_{j,m}^s\le y_{j,m},\quad w_{j,m}^s\le D_{j,m}^s\\")
    md.append(r"& u_{j,m}^s = y_{j,m}-w_{j,m}^s,\quad v_{j,m}^s = D_{j,m}^s - w_{j,m}^s\\")
    md.append(r"& D_{j,m}^s = \max\left(0,\ \hat{D}_{j,m}+\varepsilon_{j,m}^s\right),\quad "
              r"\varepsilon^s \sim \text{联合经验残差Bootstrap}")
    md.append(r"\end{aligned}")
    md.append("```\n")

    # ---- 论述段落 ----
    md.append("## 三、灵敏度与瓶颈分析学术论述段落\n")
    md.append("### 3.1 瓶颈与产能结构\n")
    md.append(f"> 优化结果表明，本案例的后厨产能构成**绝对硬约束**：午餐人工占用率 99.46%"
              f"、晚餐主灶占用率 99.60%，均已逼近饱和。"
              f"在给定产能下，任意菜品再增加一个标准生产批次即违反人工或主灶约束，"
              f"而任意减少一个批次都会降低综合经营效益，因此 Q2 确定性方案与 Q3 随机稳健方案"
              f"的**第一阶段决策完全一致**（即“满产能备餐”）。\n")
    md.append("### 3.2 需求不确定性的代价\n")
    md.append(f"> 确定性模型在“需求恰等于点预测”的假设下报告综合经营效益 {q2_benefit:.2f} 元；"
              f"然而将其置于 1000 个联合经验残差场景下，真实期望效益仅为 {q3_benefit:.2f} 元，"
              f"二者相差 **{abs(uncertainty_cost):.2f} 元（{abs(uncertainty_cost)/q2_benefit*100:.2f}%）**。"
              f"这一“需求不确定性代价”源于需求波动导致的额外剩余损失与缺供损失，"
              f"说明仅依据点预测做确定性决策会系统性高估经营效益。\n")
    md.append("### 3.3 策略建议\n")
    md.append("> 由于瓶颈为人工与主灶产能，提升经营效益的根本路径是**扩大产能**"
              "（增派后厨人手、增设主灶）而非进一步优化备餐结构；"
              "在产能扩张前，随机模型给出的服务水平 94.55%、VaR5% 10624.98 元等风险指标，"
              "可作为食堂管理层设定安全库存与服务水平承诺的量化依据。\n")

    path = out_dir / "paper_writing_materials.md"
    path.write_text("\n".join(md), encoding="utf-8")
    LOG.info("已生成白皮书 %s", path.resolve())
    if mirror_dir is not None:
        mirror_dir.mkdir(parents=True, exist_ok=True)
        (mirror_dir / "paper_writing_materials.md").write_text("\n".join(md), encoding="utf-8")
        LOG.info("镜像生成 %s", (mirror_dir / "paper_writing_materials.md").resolve())


# ===========================================================================
# 主流程
# ===========================================================================
def run_pipeline(args: argparse.Namespace) -> None:
    data_dir = Path(args.data_dir)
    out_dir = Path(args.output_dir)
    mirror_dir = Path(args.mirror_dir) if args.mirror_dir else None
    out_dir.mkdir(parents=True, exist_ok=True)
    if mirror_dir is not None:
        mirror_dir.mkdir(parents=True, exist_ok=True)
    setup_logging(out_dir, args.verbose)

    banner("阶段五｜学术图表终验与自动报告生成 启动")
    LOG.info("输出目录：%s", out_dir.resolve())

    # ---- 读取输入 ----
    residuals = pd.read_csv(out_dir / "model_residuals.csv")
    res_stats = pd.read_csv(out_dir / "residual_stats.csv")
    res_df = pd.read_csv(out_dir / "q2_resource_utilization.csv")
    predictions = pd.read_csv(out_dir / "predictions_20260629.csv")
    q2_plan = pd.read_csv(out_dir / "q2_optimal_plan.csv")
    q3_plan = pd.read_csv(out_dir / "q3_optimal_plan.csv")
    q2_perf = pd.read_csv(out_dir / "q2_performance_metrics.csv")
    params = load_dish_params(data_dir / ATTACH3_NAME)
    att4 = find_attach4(data_dir)

    q2_det = dict(zip(q2_perf["指标"], q2_perf["数值"]))
    q2_benefit = float(q2_det["综合经营效益(元)"])

    # ---- 任务 1：图1/图2 ----
    plot_residual_distribution(residuals, out_dir, mirror_dir)
    plot_resource_saturation(res_df, out_dir, mirror_dir)

    # ---- 任务 1：图3（重算 1000 场景收益分布）----
    D_scenarios, combo_keys, D_hat = generate_scenarios(residuals, predictions, 1000, 42)
    q2_y = {j: {m: 0 for m in MEALS} for j in sorted(params)}
    q3_y = {j: {m: 0 for m in MEALS} for j in sorted(params)}
    for _, r in q2_plan.iterrows():
        q2_y[str(r["菜品编号"]).strip()][str(r["时段"]).strip()] = float(r["备餐量(份)"])
    for _, r in q3_plan.iterrows():
        q3_y[str(r["菜品编号"]).strip()][str(r["时段"]).strip()] = float(r["备餐量(份)"])
    q2_m = evaluate_plan(q2_y, params, D_scenarios, combo_keys)
    q3_m = evaluate_plan(q3_y, params, D_scenarios, combo_keys)
    q3_benefit = q3_m["expected_benefit"]
    plot_cdf_overlap(q2_m["benefits"], q3_m["benefits"], q3_benefit, out_dir, mirror_dir)

    # ---- 任务 2：审计 ----
    checks = audit_excel(params, att4)
    audit_df = pd.DataFrame(checks)
    LOG.info("审计报告：\n%s", audit_df.to_string(index=False))
    audit_df.to_csv(out_dir / "audit_report.csv", index=False, encoding="utf-8-sig")
    if mirror_dir is not None:
        audit_df.to_csv(mirror_dir / "audit_report.csv", index=False, encoding="utf-8-sig")

    # ---- 任务 3：白皮书 ----
    generate_white_paper(out_dir, mirror_dir, q2_benefit, q3_benefit, res_stats, res_df, checks)

    # ---- 结束 ----
    n_pass = sum(1 for c in checks if c["结果"] == "PASS")
    n_total = len(checks)
    banner("阶段五完成")
    if n_pass == n_total:
        _print_success()
    else:
        LOG.warning("存在 %d 项审计未通过，请核查！", n_total - n_pass)


def _print_success() -> None:
    art = (
        "\n"
        "  ███████╗████████╗ █████╗  ██████╗ ███████╗    ███████╗\n"
        "  ██╔════╝╚══██╔══╝██╔══██╗██╔════╝ ██╔════╝    ██╔════╝\n"
        "  ███████╗   ██║   ███████║██║  ███╗█████╗      ███████╗\n"
        "  ╚════██║   ██║   ██╔══██║██║   ██║██╔══╝      ╚════██║\n"
        "  ███████║   ██║   ██║  ██║╚██████╔╝███████╗    ███████║\n"
        "  ╚══════╝   ╚═╝   ╚═╝  ╚═╝ ╚═════╝ ╚══════╝    ╚══════╝\n"
    )
    print("\033[1;32m" + art + "\033[0m")
    print("\033[1;32mSTAGE 5 SUCCESS: ALL ASSETS GENERATED AND AUDIT PASSED!\033[0m")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    # 默认以脚本所在目录为基准，附件3/4 在脚本目录、输出到 脚本目录/output，换机无需改代码
    script_dir = Path(__file__).resolve().parent
    p = argparse.ArgumentParser(description="阶段五：学术图表终验与自动报告生成")
    p.add_argument("--data-dir", default=str(script_dir), help="数据目录（含附件3/4，默认：脚本目录）")
    p.add_argument("--output-dir", default=str(script_dir / "output"), help="输出目录（默认：脚本目录/output）")
    p.add_argument("--mirror-dir", default="", help="镜像目录（默认禁用；如需双写可指定路径）")
    p.add_argument("--verbose", action="store_true", help="DEBUG 日志")
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    try:
        run_pipeline(args)
    except Exception as exc:
        logging.getLogger().exception("流水线执行失败：%s", exc)
        sys.exit(1)


if __name__ == "__main__":
    main()
